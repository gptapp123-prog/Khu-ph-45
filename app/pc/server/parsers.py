from __future__ import annotations
import io,csv,re,unicodedata,hashlib,lzma,tempfile,threading,atexit,os
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image,ImageOps,ImageEnhance,ImageFilter,ImageStat,ImageMath

VIETNAMESE_MODEL=Path(__file__).resolve().parents[1]/'ocr_models'/'vie.traineddata'
VIETNAMESE_MODEL_SHA256='79df64caf7bcfb2a27df5042ecb6121e196eada34da774956995747636d5bfa1'
ENGLISH_MODEL_SHA256='7d4322bd2a7749724879683fc3912cb542f19906c83bcc1a52132556427170b2'
ENGLISH_ARCHIVE_SHA256='6e287ef6ff5dac5ef95295d76f685b6695ac57a6adcdbc23956627594ffdcc5c'
_MODEL_LOCK=threading.Lock()
_MODEL_DIR=None
_MODEL_TEMP=None

def ocr_model_dir():
    global _MODEL_DIR,_MODEL_TEMP
    with _MODEL_LOCK:
        if _MODEL_DIR:return _MODEL_DIR
        if not VIETNAMESE_MODEL.is_file() or hashlib.sha256(VIETNAMESE_MODEL.read_bytes()).hexdigest()!=VIETNAMESE_MODEL_SHA256:
            raise ValueError('Thiếu hoặc sai bộ nhận dạng tiếng Việt đi kèm ứng dụng PC.')
        parts=[VIETNAMESE_MODEL.with_name('eng.traineddata.xz.part'+str(i).zfill(2)) for i in range(3)]
        if not all(p.is_file() for p in parts):raise ValueError('Thiếu bộ nhận dạng tiếng Anh đi kèm ứng dụng PC.')
        packed=b''.join(p.read_bytes() for p in parts)
        if hashlib.sha256(packed).hexdigest()!=ENGLISH_ARCHIVE_SHA256:
            raise ValueError('Bộ nhận dạng tiếng Anh không đạt kiểm tra toàn vẹn.')
        english=lzma.decompress(packed)
        if hashlib.sha256(english).hexdigest()!=ENGLISH_MODEL_SHA256:
            raise ValueError('Bộ nhận dạng tiếng Anh không đạt kiểm tra toàn vẹn.')
        tmp=tempfile.TemporaryDirectory(prefix='kpho-ocr-model-')
        path=Path(tmp.name)
        (path/'vie.traineddata').write_bytes(VIETNAMESE_MODEL.read_bytes())
        (path/'eng.traineddata').write_bytes(english)
        _MODEL_TEMP=tmp;_MODEL_DIR=path
        atexit.register(tmp.cleanup)
        return path

def ocr_image(image,requested_fields=None,mode='auto'):
    """Recognize the page, then reread selected headings in their image regions."""
    model_dir=ocr_model_dir()
    import pytesseract
    # pytesseract's Windows argument splitting can forward literal quotes in
    # --tessdata-dir. Tesseract then looks for `"C:\..."/vie.traineddata`.
    # Set once to the verified model directory inherited by every OCR subprocess.
    os.environ['TESSDATA_PREFIX']=str(model_dir)
    config=''
    try:
        prepared=prepare_ocr_image(image,force=(mode=='enhanced'))
        text=pytesseract.image_to_string(prepared,lang='vie+eng',config=config)
        if requested_fields and prepared.width>=500 and prepared.height>=300:
            text=ocr_selected_regions(prepared,text,requested_fields,pytesseract,config)
            if mode=='auto':
                from .ocr_extraction import extract_configured
                missing=[field for field in requested_fields if not extract_configured(text,[field])[0][field['column']]]
                if missing:
                    enhanced=prepare_ocr_image(image,force=True)
                    reread=pytesseract.image_to_string(enhanced,lang='vie+eng',config=config)
                    additions=[]
                    for field in missing:
                        values,warnings=extract_configured(reread,[field])
                        if values[field['column']] and not any('nhiều giá trị' in w for w in warnings):
                            additions.append(field['label']+': '+values[field['column']])
                    if additions:text+='\n[OCR tự cân sáng – cần đối chiếu]\n'+'\n'.join(additions)
        return text
    except pytesseract.TesseractNotFoundError as exc:
        raise ValueError('PC chưa cài chương trình Tesseract OCR.') from exc
    except pytesseract.TesseractError as exc:
        raise ValueError('Tesseract OCR không đọc được ảnh; kiểm tra cài đặt trên PC.') from exc

def prepare_ocr_image(image,force=False):
    """Correct EXIF orientation, lift dark/flat photos and enlarge small text in RAM."""
    image=ImageOps.exif_transpose(image).convert('RGB')
    if image.width*image.height>20_000_000:raise ValueError('Ảnh quá lớn để xử lý OCR.')
    gray=ImageOps.grayscale(image)
    stat=ImageStat.Stat(gray)
    if force and min(image.size)>=300:
        # Phone photos often have a bright side and a shadowed side. Estimate
        # the slowly changing paper background, retaining the dark glyphs.
        gray=ImageOps.grayscale(image)
        radius=min(60,max(15,min(image.size)//25))
        background=gray.filter(ImageFilter.GaussianBlur(radius))
        image=ImageMath.lambda_eval(
            lambda a:(a['paper']*230)/(a['background']+1),
            paper=gray.convert('F'),background=background.convert('F')
        ).convert('L').convert('RGB')
    elif force or stat.mean[0]<90 or stat.stddev[0]<25:
        image=ImageOps.autocontrast(image,cutoff=1)
        if stat.mean[0]<65:image=ImageEnhance.Brightness(image).enhance(1.2)
        if force:image=ImageEnhance.Contrast(image).enhance(1.2)
    target=1200 if force else 900
    if min(image.size)<target and min(image.size)>=120:
        factor=min(2.0,target/min(image.size))
        if image.width*image.height*factor*factor<=20_000_000:
            image=image.resize((round(image.width*factor),round(image.height*factor)),Image.Resampling.LANCZOS if hasattr(Image,'Resampling') else Image.LANCZOS)
            image=image.filter(ImageFilter.UnsharpMask(radius=1,percent=110,threshold=3))
    return image

def ocr_selected_regions(image,text,fields,pytesseract,config):
    """Use word coordinates to crop a heading and its displaced right-hand value.

    Only missing fields are reconsidered. A conflicting or unverified number is
    never promoted into the field; the original page text remains available.
    """
    from .ocr_extraction import extract_configured,normalized
    missing=[f for f in fields if not extract_configured(text,[f])[0][f['column']]]
    if not missing:return text
    data=pytesseract.image_to_data(image,lang='vie+eng',config=config,output_type=pytesseract.Output.DICT)
    tokens=[]
    for i,word in enumerate(data['text']):
        clean=normalized(word).strip(' :.,;_–—')
        if clean:tokens.append((clean,i))
    additions=[]
    for field in missing[:30]:
        candidates=[]
        for alias in [field['label']]+field['aliases']:
            words=[normalized(w).strip(' :.,;_–—') for w in alias.split()]
            for start in range(len(tokens)-len(words)+1):
                group=tokens[start:start+len(words)]
                if [t[0] for t in group]!=words:continue
                indices=[t[1] for t in group]
                if max(data['top'][j] for j in indices)-min(data['top'][j] for j in indices)>40:continue
                left=min(data['left'][j] for j in indices)
                top=min(data['top'][j] for j in indices)
                height=max(data['height'][j] for j in indices)
                # Include the whole line: certificate suffixes and chassis
                # numbers can sit at the far right of a wide landscape photo.
                x0=max(0,left-25);x1=image.width
                y0=max(0,top-max(47,int(height*1.5)));y1=min(image.height,top+max(55,int(height*1.8)))
                crop=image.crop((x0,y0,x1,y1))
                factor=min(3, max(1, 3600/crop.width))
                crop=crop.resize((round(crop.width*factor),round(crop.height*factor)))
                crop_text=pytesseract.image_to_string(crop,lang='vie+eng',config=config+' --psm 6')
                value=extract_configured(crop_text,[field])[0][field['column']]
                if not value:continue
                value=re.sub(r'(?<=\d)\s+[_|]\s+(?=/)',' ',value)
                digits=re.findall(r'\d{4,}',value)
                if digits:
                    nearby=[str(data['text'][j]).strip('.,;') for j in range(len(data['text']))
                            if abs(data['top'][j]-top)<=70 and data['left'][j]>=left and float(data['conf'][j])>=70]
                    if not any(number in nearby for number in digits):continue
                candidates.append(value)
        unique=list(dict.fromkeys(candidates))
        if len(unique)==1:additions.append(field['label']+': '+unique[0])
    return text+('\n[OCR đọc lại vùng ảnh – cần đối chiếu]\n'+'\n'.join(additions) if additions else '')

KEYWORDS={'cccd':['cccd','can cuoc','so dinh danh','so cccd'],
'full_name':['ho va ten','ho ten','ten day du','ten'], 'dob':['ngay sinh','sinh ngay'],
'gender':['gioi tinh'], 'hometown':['que quan'],
'permanent_address':['thuong tru','dia chi','noi thuong tru'],
'team_id':['to dan pho','to','to so'], 'phone':['so dien thoai','sdt','dien thoai'],
'occupation':['nghe nghiep'], 'health_insurance':['so the bhyt','so the bao hiem y te','bhyt','bao hiem y te'],
'note':['ghi chu']}

def strip_accents(v):
    return ''.join(c for c in unicodedata.normalize('NFD',str(v).lower().replace('đ','d')) if unicodedata.category(c)!='Mn').strip()
def validate_headers(headers):
    cleaned=[str(h).strip() for h in headers]
    if not cleaned or not any(cleaned):raise ValueError('Tệp không có tiêu đề cột.')
    normalized=[strip_accents(h) for h in cleaned]
    if any(not h for h in normalized):raise ValueError('Tiêu đề cột không được để trống.')
    if len(normalized)!=len(set(normalized)):raise ValueError('Tiêu đề cột bị trùng.')
    return cleaned

def automap(headers):
    mapped={}
    for h in headers:
        s=strip_accents(h)
        for k,variants in KEYWORDS.items():
            if s in variants: mapped[h]=k;break
    return mapped

OCR_LABELS={'full_name':('ho va ten','ho ten','full name'),
            'dob':('ngay thang nam sinh','ngay sinh','date of birth'),
            'gender':('gioi tinh','sex'),
            'hometown':('que quan','place of origin'),
            'permanent_address':('noi thuong tru','thuong tru','place of residence'),
            'cccd':('so can cuoc','so dinh danh ca nhan','so cccd','cccd','identity number')}
OCR_HEADERS={'full_name':'Họ tên OCR','dob':'Ngày sinh OCR','gender':'Giới tính OCR',
             'hometown':'Quê quán OCR','permanent_address':'Thường trú OCR','cccd':'Số CCCD nhận dạng'}

def extract_ocr_fields(content):
    """Extract only clearly labelled values; ambiguous values require review."""
    found={k:[] for k in OCR_HEADERS}
    lines=content.splitlines()
    for index,line in enumerate(lines):
        line=line.strip()
        if not line:continue
        normalized=strip_accents(line)
        matches=[]
        for field,labels in OCR_LABELS.items():
            for label in labels:
                match=re.match(r'^'+re.escape(label)+r'\s*[:：-]\s*(.*)$',normalized,re.I)
                if match:
                    value=line[match.start(1):].strip()
                    if not value and index+1<len(lines):
                        following=lines[index+1].strip()
                        next_label=strip_accents(following).split(':',1)[0].strip()
                        if next_label not in {name for group in OCR_LABELS.values() for name in group}:
                            value=following
                    matches.append((field,value))
                    break
        if len(matches)==1:found[matches[0][0]].append(matches[0][1])
    candidates=set(re.findall(r'(?<!\d)\d{12}(?!\d)',content))
    result={};warnings=[]
    for field,values in found.items():
        unique={v.strip() for v in values if v.strip()}
        if field=='cccd':
            unique={v for value in unique for v in re.findall(r'(?<!\d)\d{12}(?!\d)',value)}
            if len(candidates)>1:unique=set()
            if not unique and len(candidates)==1:unique=candidates
        if field=='dob':
            unique={v for v in unique if re.fullmatch(r'\d{2}[-/]\d{2}[-/]\d{4}',v)}
        if len(unique)==1:result[field]=next(iter(unique))
        elif len(unique)>1:warnings.append('Có nhiều giá trị cho '+OCR_HEADERS[field]+'; cần đối chiếu ảnh.')
    if len(candidates)>1:warnings.append('Có nhiều số 12 chữ số; không tự chọn CCCD.')
    return result,warnings

def review_result(texts,source):
    content='\n'.join(texts)
    extracted,warnings=extract_ocr_fields(content)
    if '[OCR đọc lại vùng ảnh – cần đối chiếu]' in content:
        warnings.append('Có giá trị đọc lại từ vùng ảnh. Phải đối chiếu từng ký tự với ảnh gốc trước khi xác nhận.')
    if '[OCR tự cân sáng – cần đối chiếu]' in content:
        warnings.append('Có giá trị từ lượt tăng sáng/tương phản. Phải đối chiếu từng ký tự với ảnh gốc trước khi xác nhận.')
    if '[OCR bổ sung từ trang PDF – cần đối chiếu]' in content:
        warnings.append('Trang PDF có lớp chữ không đầy đủ; giá trị bổ sung từ ảnh scan cần đối chiếu với trang gốc.')
    headers=list(OCR_HEADERS.values())
    row={label:extracted.get(field,'') for field,label in OCR_HEADERS.items()}
    items=[{'source':label,'label':label,'value':row[label],'suggested':field}
           for field,label in OCR_HEADERS.items()]
    seen={strip_accents(label) for label in OCR_HEADERS.values()}
    for line in content.splitlines():
        match=re.match(r'^\s*([^:：]{2,80})\s*[:：]\s*(.{1,1000})\s*$',line)
        if not match:continue
        label,value=match.group(1).strip(),match.group(2).strip()
        if not label or not value:continue
        matched=next((field for field,labels in OCR_LABELS.items()
                      if strip_accents(label) in labels),None)
        if matched:continue
        key=strip_accents(label)
        if key in seen:
            warnings.append('Nhãn OCR lặp lại: '+label+'; kiểm tra văn bản gốc.')
            continue
        seen.add(key);headers.append(label);row[label]=value
        items.append({'source':label,'label':label,'value':value,'suggested':automap([label]).get(label,'')})
        if len(items)>=100:
            warnings.append('Chỉ hiển thị 100 trường OCR đầu tiên; hãy tách tài liệu.')
            break
    rows=[row]
    result={'headers':headers,'preview':rows,'rows':rows,'count':1,
            'mapping':{label:field for field,label in OCR_HEADERS.items() if extracted.get(field)},
            'ocr_review':True,'ocr_items':items,'raw_text':content[:12000],
            'warnings':['Nội dung '+source+' chỉ là gợi ý. Đối chiếu bản gốc và chọn từng trường trước khi ghi.']+warnings}
    if len(texts)>1:result['warnings'].append('Tài liệu nhiều trang: kiểm tra thông tin thuộc cùng một người trước khi nhập.')
    return result

def render_pdf_page(raw,page_number,requested_fields=None,ocr_mode='auto'):
    """Render a scan entirely in RAM when pypdf cannot extract its image."""
    try:import pypdfium2 as pdfium
    except ImportError as exc:raise ValueError('PDF scan cần cài pypdfium2 trên PC để kết xuất toàn trang.') from exc
    document=pdfium.PdfDocument(raw)
    try:
        page=document[page_number]
        try:
            width,height=page.get_size()
            scale=2.0
            if width*height*scale*scale>20_000_000:
                raise ValueError('Trang PDF quá lớn để OCR an toàn.')
            bitmap=page.render(scale=scale)
            try:
                with bitmap.to_pil() as image:
                    return ocr_image(image,requested_fields,ocr_mode)
            finally:bitmap.close()
        finally:page.close()
    finally:document.close()

def parse_file(raw:bytes,filename:str,requested_fields=None,ocr_mode='auto',sheet_name=None):
    if len(raw)>10*1024*1024:raise ValueError('Giới hạn tệp nhập: 10 MB.')
    ext=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    rows=[]
    if ext in ('xlsx','xlsm'):
        # Normal worksheet mode is intentional: read-only worksheets do not expose
        # merged-cell ranges, while rejecting merged input is required to prevent
        # columns from being shifted into the wrong resident fields. The upload is
        # already capped at 10 MB and row ingestion at 5,000.
        wb=load_workbook(io.BytesIO(raw),read_only=False,data_only=False)
        headers=[]
        try:
            if sheet_name is not None and (sheet_name not in wb.sheetnames or wb[sheet_name].sheet_state!='visible'):
                raise ValueError('Sheet được chọn không tồn tại hoặc đang ẩn.')
            for sh in wb.worksheets:
                if sheet_name is not None and sh.title!=sheet_name:continue
                # Hidden sheets are commonly lookup/configuration data. Importing them
                # silently can disclose or duplicate records, so only visible sheets
                # participate in a resident import.
                if sh.sheet_state!='visible':continue
                if sh.merged_cells.ranges:
                    raise ValueError(f'Sheet "{sh.title}" có ô gộp. Hãy bỏ gộp ô trước khi nhập.')
                it=sh.iter_rows()
                first=next(it,None)
                if not first:continue
                if not any(c.value is not None and str(c.value).strip() for c in first):
                    # Do not guess which later row is the header: that can map data to
                    # the wrong fields. Empty visible sheets are ignored; populated
                    # sheets must put the header on row 1.
                    if any(any(c.value is not None and str(c.value).strip() for c in row)
                           for row in it):
                        raise ValueError(f'Sheet "{sh.title}" phải đặt tiêu đề ở dòng đầu tiên.')
                    continue
                if any(c.data_type=='f' for c in first):raise ValueError('Tiêu đề Excel không được chứa công thức.')
                if any(c.value is None or not str(c.value).strip() for c in first):
                    raise ValueError(f'Sheet "{sh.title}" có cột không có tiêu đề.')
                sheet_headers=validate_headers([str(c.value).strip() if c.value is not None else f'Cột {i+1}' for i,c in enumerate(first)])
                if headers and [strip_accents(x) for x in sheet_headers]!=[strip_accents(x) for x in headers]:
                    raise ValueError('Các sheet Excel phải có cùng tiêu đề và thứ tự cột.')
                if not headers:headers=sheet_headers
                for cells in it:
                    if any(c.data_type=='f' for c in cells):raise ValueError('Excel chứa công thức. Hãy chuyển công thức thành giá trị trước khi nhập.')
                    values=[c.value for c in cells]
                    if any(v is not None and str(v).strip() for v in values):
                        if any(v is not None and str(v).strip() for v in values[len(headers):]):
                            raise ValueError(f'Sheet "{sh.title}" có dữ liệu ngoài các cột tiêu đề.')
                        rows.append({headers[i]: '' if i>=len(values) or values[i] is None else str(values[i]).strip() for i in range(len(headers))})
                    if len(rows)>5000:raise ValueError('Tối đa 5.000 dòng/lần nhập.')
            if not headers:raise ValueError('Excel trống.')
        finally:wb.close()
    elif ext=='csv':
        content=raw.decode('utf-8-sig')
        first=content[:1024]
        try:dialect=csv.Sniffer().sniff(first,delimiters=',;\t')
        except csv.Error:dialect=csv.excel
        reader=csv.DictReader(io.StringIO(content),dialect=dialect)
        headers=validate_headers(reader.fieldnames or [])
        reader.fieldnames=headers
        for row in reader:
            if len(rows)>=5000:raise ValueError('Tối đa 5.000 dòng/lần nhập.')
            if None in row:raise ValueError('CSV có dòng nhiều cột hơn tiêu đề.')
            rows.append({str(k or ''):str(v or '') for k,v in row.items()})
    elif ext=='pdf':
        pdf=PdfReader(io.BytesIO(raw))
        if len(pdf.pages)>50:raise ValueError('PDF tối đa 50 trang/lần nhập.')
        texts=[]
        for page_number,page in enumerate(pdf.pages):
            content=(page.extract_text() or '').strip()
            if content and requested_fields:
                # A scanned form may have a searchable watermark or page number
                # while its actual fields remain pixels. Check the requested
                # columns rather than treating any extracted text as complete.
                from .ocr_extraction import extract_configured
                existing,_=extract_configured(content,requested_fields)
                missing=[f for f in requested_fields if not existing[f['column']]]
                if missing:
                    scanned=render_pdf_page(raw,page_number,missing,ocr_mode)
                    recovered,_=extract_configured(scanned,missing)
                    additions=[f['label']+': '+recovered[f['column']] for f in missing if recovered[f['column']]]
                    if additions:
                        content+='\n[OCR bổ sung từ trang PDF – cần đối chiếu]\n'+'\n'.join(additions)
            if not content:
                images=list(page.images)
                if not images:
                    texts.append(render_pdf_page(raw,page_number,requested_fields,ocr_mode))
                    continue
                if len(images)>10:raise ValueError('Trang PDF có quá 10 ảnh; hãy tách tài liệu để tránh bỏ sót.')
                page_text=[]
                for item in images[:10]:
                    with Image.open(io.BytesIO(item.data)) as img:
                        if img.width*img.height>20_000_000:raise ValueError('Ảnh trong PDF quá lớn.')
                        page_text.append(ocr_image(img,requested_fields,ocr_mode))
                content='\n'.join(page_text)
            texts.append(content)
        if not any(x.strip() for x in texts):raise ValueError('Không đọc được nội dung PDF.')
        return review_result(texts,'PDF')
    elif ext in ('png','jpg','jpeg','webp'):
        with Image.open(io.BytesIO(raw)) as img:
            if img.width*img.height>20_000_000:raise ValueError('Ảnh quá lớn.')
            content=ocr_image(img,requested_fields,ocr_mode)
        return review_result([content],'OCR ảnh')
    else:raise ValueError('Định dạng hỗ trợ: XLSX, CSV, PDF, PNG, JPG, WEBP.')
    result={'headers':headers,'preview':rows[:10],'rows':rows,'mapping':automap(headers),'count':len(rows)}
    return result
