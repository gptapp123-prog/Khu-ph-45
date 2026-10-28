"""QR is decoded locally and mapped only to configured destination columns."""
import io,json
import qrcode
from PIL import Image
from server.qr_extraction import qr_values,decode_qr_file
from test_foundation_accounts import api  # noqa: F401

FIELDS=[{'label':'Họ và tên','column':'Họ và tên','aliases':['Họ tên'],'type':'text'},
        {'label':'Số hiệu bằng','column':'Số hiệu bằng','aliases':[],'type':'identifier'},
        {'label':'Nội dung QR','column':'Nội dung QR','aliases':[],'type':'text'}]

def qr_png(value):
    image=qrcode.make(value).convert('RGB')
    output=io.BytesIO();image.save(output,format='PNG')
    return output.getvalue()

def test_decode_real_qr_and_export_only_selected_columns(api):
    client,headers,_=api
    template=client.post('/api/ocr/templates',headers=headers,json={'name':'Văn bằng QR','fields':FIELDS})
    assert template.status_code==200,template.text
    tid=template.json()['id']
    payload='Họ và tên: Nguyễn Văn A\nSố hiệu bằng: 12345/PT'
    raw=qr_png(payload)
    decoded=client.post('/api/qr/decode?template_id='+tid,headers=headers,files={'file':('ma.png',raw,'image/png')})
    assert decoded.status_code==200,decoded.text
    body=decoded.json()
    assert body['codes']==[{'page':1,'text':payload}]
    assert body['values']['Họ và tên']=='Nguyễn Văn A'
    assert body['values']['Số hiệu bằng']=='12345/PT'
    assert body['values']['Nội dung QR']==payload
    row=body['values']
    export=client.post('/api/ocr/export',headers=headers,json={'template_id':tid,'filename':'qr-result','rows':[row]})
    assert export.status_code==200
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(export.content));ws=wb.active
    assert [cell.value for cell in ws[1]]==list(row)
    assert ws['B2'].value=='12345/PT'
    wb.close()

def test_unlabelled_qr_is_never_guessed_as_personal_fields():
    payload='079123456789|NGUYEN VAN A|19900101|Nam'
    values,_=qr_values([{'page':1,'text':payload}],FIELDS)
    assert values=={'Họ và tên':'','Số hiệu bằng':'','Nội dung QR':payload}
    obj=json.dumps({'Họ tên':'Nguyễn Văn A','Số hiệu bằng':'12345/PT'},ensure_ascii=False)
    values,_=qr_values([{'page':1,'text':obj}],FIELDS[:2])
    assert values=={'Họ và tên':'Nguyễn Văn A','Số hiệu bằng':'12345/PT'}

def test_multiple_qr_with_conflicting_field_stays_blank():
    values,warnings=qr_values([{'page':1,'text':'Số hiệu bằng: 12345/PT'},
                               {'page':2,'text':'Số hiệu bằng: 12346/PT'}],FIELDS[:2])
    assert values['Số hiệu bằng']==''
    assert any('nhiều QR' in warning for warning in warnings)

def test_qr_image_format_limits_and_missing_template(api):
    client,headers,_=api
    assert client.post('/api/qr/decode?template_id=missing',headers=headers,files={'file':('ma.png',qr_png('A'),'image/png')}).status_code==404
    assert decode_qr_file(qr_png('Địa chỉ: Long An'),'ma.png')[0]['text']=='Địa chỉ: Long An'
    from pytest import raises
    with raises(ValueError,match='hỗ trợ'):
        decode_qr_file(b'not an image','ma.txt')

def test_qr_in_scanned_pdf_page():
    image=Image.open(io.BytesIO(qr_png('Số hiệu bằng: 12345/PT')))
    output=io.BytesIO();image.save(output,format='PDF',resolution=150)
    assert decode_qr_file(output.getvalue(),'scan.pdf')==[{'page':1,'text':'Số hiệu bằng: 12345/PT'}]
