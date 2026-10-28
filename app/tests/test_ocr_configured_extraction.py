"""Configured headings control both extracted values and Excel columns."""
import io
from unittest.mock import patch
from openpyxl import load_workbook
from PIL import Image
from server.ocr_extraction import extract_configured
from server.parsers import parse_file
from server.parsers import ocr_image,prepare_ocr_image
from test_foundation_accounts import api  # noqa: F401 - shared vault fixture


def fields():
    return [
        {'label':'Họ và tên','aliases':['Họ tên'],'column':'Họ và tên','type':'text'},
        {'label':'Số hiệu bằng','aliases':[],'column':'Số hiệu bằng','type':'text'},
        {'label':'Học sinh trường','aliases':['Trường'],'column':'Trường','type':'text'},
        {'label':'Tại hội đồng thi','aliases':['Hội đồng thi'],'column':'Hội đồng thi','type':'text'},
    ]


def test_full_values_suffix_and_multiple_document_headings():
    text='Họ và tên: Nguyễn Văn A\nSố hiệu bằng: 12345/ptxn\nHọc sinh trường: THPT Thủ Đức\nTại hội đồng thi: THPT Tam Phú\nNghề nghiệp: Kế toán'
    result,warnings=extract_configured(text,fields())
    assert result=={'Họ và tên':'Nguyễn Văn A','Số hiệu bằng':'12345/ptxn','Trường':'THPT Thủ Đức','Hội đồng thi':'THPT Tam Phú'}
    assert warnings==[] and 'Nghề nghiệp' not in result


def test_dotted_leaders_multiline_and_adjacent_heading():
    text='Số hiệu bằng:..... 2244863 /PT.\nHọ và tên:\nNguyễn Văn A\nHọc sinh trường: THPT Thủ Đức\nTại hội đồng thi: THPT Tam Phú'
    result,_=extract_configured(text,fields())
    assert result['Số hiệu bằng']=='2244863 /PT.'
    assert result['Họ và tên']=='Nguyễn Văn A'
    only=[{'label':'Ngày sinh','aliases':[],'column':'Ngày sinh','type':'text'}]
    assert extract_configured('Ngày sinh: 18/07/1988 Giới tính: Nam',only)[0]['Ngày sinh']=='18/07/1988'
    both=only+[{'label':'Giới tính','aliases':[],'column':'Giới tính','type':'text'}]
    assert extract_configured('Ngày sinh: 18/07/1988 Giới tính: Nam',both)[0]=={'Ngày sinh':'18/07/1988','Giới tính':'Nam'}


def test_heading_mentioned_in_prose_is_not_extracted():
    f=[{'label':'Số khung','aliases':[],'column':'Số khung','type':'identifier'}]
    text='Ghi chú: cần kiểm tra Số khung: RL4JW3EM3M0000000 trên bản gốc'
    values,warnings=extract_configured(text,f)
    assert values['Số khung']==''
    assert any('không tìm thấy' in warning for warning in warnings)
    same_row='Nhãn hiệu (Brand): TOYOTA Số loại (Model code): INNOVA'
    f=[{'label':'Nhãn hiệu','aliases':[],'column':'Nhãn hiệu','type':'text'},
       {'label':'Số loại','aliases':[],'column':'Số loại','type':'text'}]
    values,_=extract_configured(same_row,f)
    assert values=={'Nhãn hiệu':'TOYOTA','Số loại':'INNOVA'}


def test_conflict_and_missing_stay_blank():
    result,warnings=extract_configured('Số hiệu bằng: 12345/PT\nSố hiệu bằng: 12346/PT\nHọ và tên:\nNgày sinh: 01/01/2000',fields())
    assert result['Số hiệu bằng']=='' and result['Họ và tên']==''
    assert any('nhiều giá trị' in w for w in warnings)


def test_ocr_never_invents_value_from_following_unrelated_heading():
    text='Số hiệu bằng:\n\nv Họ và tên: Nguyễn Văn A\nHọc sinh trường: , : THPT Thủ Đức'
    result,warnings=extract_configured(text,fields())
    assert result['Số hiệu bằng']==''
    assert result['Họ và tên']=='Nguyễn Văn A'
    assert result['Trường']=='THPT Thủ Đức'
    assert any('Số hiệu bằng' in warning for warning in warnings)


def test_ocr_uses_bundled_vietnamese_model_and_checks_integrity():
    from unittest.mock import patch
    import pytest
    with patch('pytesseract.image_to_string',return_value='Học sinh trường: THPT Thủ Đức') as mock:
        assert 'THPT Thủ Đức' in ocr_image(Image.new('RGB',(30,30),'white'))
        assert mock.call_args.kwargs['lang']=='vie+eng'
        assert '--tessdata-dir' in mock.call_args.kwargs['config']
    with patch('server.parsers._MODEL_DIR',None),patch('server.parsers.VIETNAMESE_MODEL') as path:
        path.is_file.return_value=False
        with pytest.raises(ValueError,match='Thiếu hoặc sai bộ nhận dạng'):
            ocr_image(Image.new('RGB',(30,30),'white'))


def test_layout_crop_recovers_displaced_identifier_and_rejects_wrong_digits():
    from unittest.mock import patch
    f=[{'label':'Số hiệu bằng','aliases':[],'column':'Số hiệu bằng','type':'text'}]
    image=Image.new('RGB',(1000,600),'white')
    data={'text':['Số','hiệu','bằng:','12345'],
          'left':[100,135,180,310],'top':[100,98,95,75],
          'height':[20,23,27,30],'conf':[95,95,95,92]}
    page='Số hiệu bằng:\nHọ và tên: Nguyễn Văn A'
    with patch('pytesseract.image_to_data',return_value=data),patch('pytesseract.image_to_string',side_effect=[page,'Số hiệu bằng: 12345 /PT.']):
        text=ocr_image(image,f)
    assert extract_configured(text,f)[0]['Số hiệu bằng']=='12345 /PT.'
    assert '[OCR đọc lại vùng ảnh' in text
    with patch('pytesseract.image_to_data',return_value=data),patch('pytesseract.image_to_string',side_effect=[page,'Số hiệu bằng: 99999 /PT.','']):
        text=ocr_image(image,f)
    assert extract_configured(text,f)[0]['Số hiệu bằng']==''


def test_auto_enhanced_retry_only_fills_missing_configured_field():
    from unittest.mock import patch
    f=[{'label':'Số khung','aliases':[],'column':'Số khung','type':'identifier','expected_length':17}]
    data={'text':[],'left':[],'top':[],'height':[],'conf':[]}
    primary='Số khung: RL4W3EM3M3375080\nNhãn khác: Bỏ qua'
    enhanced='Số khung: RL4JW3EM3M3375080\nNhãn khác: Bỏ qua'
    with patch('pytesseract.image_to_data',return_value=data),patch('pytesseract.image_to_string',side_effect=[primary,enhanced]):
        text=ocr_image(Image.new('RGB',(1000,600),'white'),f)
    values,warnings=extract_configured(text,f)
    assert values['Số khung']=='RL4JW3EM3M3375080'
    assert 'Nhãn khác' not in values
    assert '[OCR tự cân sáng' in text


def test_bilingual_labels_multiline_dates_and_identifier_length():
    configured=[
        {'label':'Tên chủ xe','aliases':[],'column':'Chủ xe','type':'text'},
        {'label':'Địa chỉ','aliases':[],'column':'Địa chỉ','type':'multiline'},
        {'label':'Số khung','aliases':[],'column':'Số khung','type':'identifier','expected_length':17},
        {'label':'Ngày cấp','aliases':[],'column':'Ngày cấp','type':'date'},
    ]
    text="Tên chủ xe (Owner's full name):\nNGUYỄN THỊ MINH AN\nĐịa chỉ (Address): 123 Tỉnh Lộ 10, Quận Bình\nTân\nSố khung (Chassis N°): RL4JW3EM3M0000000\nNgày cấp: 27 tháng 01 năm 2021"
    values,warnings=extract_configured(text,configured)
    assert values=={'Chủ xe':'NGUYỄN THỊ MINH AN','Địa chỉ':'123 Tỉnh Lộ 10, Quận Bình Tân','Số khung':'RL4JW3EM3M0000000','Ngày cấp':'27-01-2021'}
    assert warnings==[]
    text=text.replace('RL4JW3EM3M0000000','RL4JW3EM30000000')
    values,warnings=extract_configured(text,configured)
    assert values['Số khung']=='' and any('độ dài' in w for w in warnings)


def test_auto_light_and_zoom_only_for_small_or_dark_images():
    image=Image.new('RGB',(450,300),(35,35,35))
    from PIL import ImageDraw
    ImageDraw.Draw(image).text((10,10),'TEST',fill='white')
    prepared=prepare_ocr_image(image)
    assert min(prepared.size)>=600
    assert prepared.getpixel((5,5))[0]>35
    normal=Image.new('RGB',(1400,1000),'white')
    assert prepare_ocr_image(normal).size==normal.size
    assert min(prepare_ocr_image(Image.new('RGB',(800,1000),'white'),force=True).size)>=1200


def test_enhanced_photo_flattens_uneven_paper_light():
    from PIL import ImageDraw
    image=Image.new('RGB',(1200,600))
    for x in range(image.width):
        shade=110+round(130*x/(image.width-1))
        for y in range(image.height):image.putpixel((x,y),(shade,shade,shade))
    ImageDraw.Draw(image).rectangle((500,240,700,260),fill=(20,20,20))
    enhanced=prepare_ocr_image(image,force=True)
    scale=enhanced.width/image.width
    left=enhanced.getpixel((round(100*scale),round(100*scale)))[0]
    right=enhanced.getpixel((round(1100*scale),round(100*scale)))[0]
    ink=enhanced.getpixel((round(600*scale),round(250*scale)))[0]
    assert abs(left-right)<35
    assert ink<min(left,right)-70


def test_reread_region_reaches_far_right_suffix():
    from server.parsers import ocr_selected_regions
    f=[{'label':'Số hiệu bằng','aliases':[],'column':'Số hiệu bằng','type':'identifier'}]
    image=Image.new('RGB',(1800,600),'white')
    data={'text':['Số','hiệu','bằng:','12345'],
          'left':[100,135,180,1400],'top':[100,100,100,100],
          'height':[25,25,25,25],'conf':[95,95,95,95]}
    class FakeOCR:
        class Output:DICT='dict'
        def image_to_data(self,*args,**kwargs):return data
        def image_to_string(self,crop,**kwargs):
            assert crop.width>=1400
            return 'Số hiệu bằng: 12345 /PT.'
    result=ocr_selected_regions(image,'Số hiệu bằng:\nHọ và tên: Nguyễn A',f,FakeOCR(),'')
    assert extract_configured(result,f)[0]['Số hiệu bằng']=='12345 /PT.'


def test_api_template_image_extract_and_excel_export(api):
    client,headers,_=api
    saved=client.post('/api/ocr/templates',headers=headers,json={'name':'Thông tin văn bằng','fields':fields()})
    assert saved.status_code==200,saved.text
    tid=saved.json()['id']
    image=Image.new('RGB',(60,30),'white');stream=io.BytesIO();image.save(stream,format='PNG')
    text='Họ và tên: Nguyễn Văn A\nSố hiệu bằng: 12345/ptxn\nHọc sinh trường: THPT Thủ Đức\nTại hội đồng thi: THPT Tam Phú\nNghề nghiệp: Kế toán'
    with patch('pytesseract.image_to_string',return_value=text):
        preview=client.post('/api/import/preview?ocr_template_id='+tid,headers=headers,files={'file':('anh.png',stream.getvalue(),'image/png')})
    assert preview.status_code==200,preview.text
    invalid_mode=client.post('/api/import/preview?ocr_mode=unsafe',headers=headers,files={'file':('anh.png',stream.getvalue(),'image/png')})
    assert invalid_mode.status_code==422
    with patch('pytesseract.image_to_string',return_value=text):
        enhanced=client.post('/api/import/preview?ocr_template_id='+tid+'&ocr_mode=enhanced',headers=headers,files={'file':('anh.png',stream.getvalue(),'image/png')})
    assert enhanced.status_code==200 and any('tăng tương phản' in x for x in enhanced.json()['warnings'])
    extracted=client.post('/api/ocr/extract',headers=headers,json={'template_id':tid,'text':preview.json()['raw_text']})
    assert extracted.status_code==200,extracted.text
    row=extracted.json()['values']
    assert row['Số hiệu bằng']=='12345/ptxn' and row['Hội đồng thi']=='THPT Tam Phú'
    assert 'Nghề nghiệp' not in row
    output=client.post('/api/ocr/export',headers=headers,json={'template_id':tid,'filename':'van_bang.xlsx','rows':[row]})
    assert output.status_code==200,output.text
    wb=load_workbook(io.BytesIO(output.content));ws=wb.active
    assert [c.value for c in ws[1]]==list(row)
    assert [c.value for c in ws[2]]==list(row.values())
    wb.close()
    invalid=client.post('/api/ocr/export',headers=headers,json={'template_id':tid,'filename':'van_bang','rows':[{**row,'Nghề nghiệp':'Kế toán'}]})
    assert invalid.status_code==422


def test_pdf_with_searchable_watermark_rereads_missing_selected_fields():
    from unittest.mock import MagicMock
    page=MagicMock();page.extract_text.return_value='Mã trang 42\nSố hiệu bằng: 12345 /PT.'
    document=MagicMock();document.pages=[page]
    scanned='Học sinh trường: THPT Thủ Đức\nTại hội đồng thi: THPT Tam Phú\nNghề nghiệp: Kế toán'
    with patch('server.parsers.PdfReader',return_value=document),patch('server.parsers.render_pdf_page',return_value=scanned) as render:
        result=parse_file(b'%PDF synthetic','scan.pdf',fields())
    row,_=extract_configured(result['raw_text'],fields())
    assert row['Số hiệu bằng']=='12345 /PT.'
    assert row['Trường']=='THPT Thủ Đức' and row['Hội đồng thi']=='THPT Tam Phú'
    assert 'Nghề nghiệp' not in row
    assert any('lớp chữ không đầy đủ' in warning for warning in result['warnings'])
    assert render.call_args.args[1]==0
    assert [f['column'] for f in render.call_args.args[2]]==['Họ và tên','Trường','Hội đồng thi']


def test_pdf_ocr_does_not_choose_conflicting_scan_values():
    from unittest.mock import MagicMock
    page=MagicMock();page.extract_text.return_value='Số hiệu bằng: 12345 /PT.\nMã trang 42'
    document=MagicMock();document.pages=[page]
    scanned='Họ và tên: Nguyễn A\nHọ và tên: Nguyễn B\nTrường: THPT Thủ Đức'
    with patch('server.parsers.PdfReader',return_value=document),patch('server.parsers.render_pdf_page',return_value=scanned):
        result=parse_file(b'%PDF synthetic','scan.pdf',fields())
    row,_=extract_configured(result['raw_text'],fields())
    assert row['Họ và tên']==''
    assert row['Trường']=='THPT Thủ Đức'
