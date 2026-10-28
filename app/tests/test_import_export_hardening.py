"""Nhập/xuất phải giữ đúng dữ liệu, quyền và an toàn bảng tính."""
import csv
import io
import json
from openpyxl import load_workbook
import pytest
from fastapi.testclient import TestClient
from server.app import app,configure_vault,SESSIONS

PW='KhoNhapXuat_2026!'
ADMIN='QuanTriNhapXuat_2026!'

@pytest.fixture
def api(tmp_path):
    SESSIONS.sessions.clear()
    configure_vault(str(tmp_path/'import_export.vault'),PW)
    with TestClient(app) as client:
        assert client.post('/api/setup',json={'username':'quantri','full_name':'Quản trị nhập xuất','password':ADMIN}).status_code==200
        login=client.post('/api/login',json={'username':'quantri','password':ADMIN})
        yield client,{'Authorization':'Bearer '+login.json()['token']}

def test_csv_semicolon_preview_commit_and_json_export(api):
    client,headers=api
    raw='Họ tên;Số CCCD;Tổ dân phố;Ghi chú\nNguyễn Thị Ánh;079123456789;3;Hồ sơ thử\n'.encode('utf-8-sig')
    preview=client.post('/api/import/preview',headers=headers,files={'file':('dan-cu.csv',raw,'text/csv')})
    assert preview.status_code==200,preview.text
    data=preview.json()
    assert data['count']==1
    assert data['mapping']['Họ tên']=='full_name'
    assert data['mapping']['Tổ dân phố']=='team_id'
    commit=client.post('/api/import/commit',headers=headers,json={
        'rows':data['rows'],'mapping':data['mapping'],
        'selected':['full_name','cccd','team_id','note']})
    assert commit.status_code==200,commit.text
    assert commit.json()=={'inserted':1,'duplicates':0,'possible_duplicates':0,
                           'warnings':[],'invalid':0,'errors':[]}
    exported=client.get('/api/export',headers=headers,params={'fields':'full_name,cccd,team_id,note','format':'json'})
    assert exported.status_code==200
    assert exported.json()[0]['full_name']=='Nguyễn Thị Ánh'

@pytest.mark.parametrize('header',[
    'Họ tên,Họ tên,Tổ\nNguyễn A,Nguyễn B,1\n',
    'Họ tên,,Tổ\nNguyễn A,x,1\n'])
def test_csv_rejects_duplicate_or_empty_headers(api,header):
    client,headers=api
    response=client.post('/api/import/preview',headers=headers,files={'file':('loi.csv',header.encode('utf-8'),'text/csv')})
    assert response.status_code==422

def test_csv_and_xlsx_exports_neutralize_formula_cells_but_json_preserves_text(api):
    client,headers=api
    payload={'full_name':'=HYPERLINK("https://invalid.example","Bấm")','team_id':1,'note':'@SUM(1+1)'}
    created=client.post('/api/residents',headers=headers,json=payload)
    assert created.status_code==200,created.text
    csv_response=client.get('/api/export',headers=headers,params={'fields':'full_name,note','format':'csv'})
    rows=list(csv.reader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
    assert rows[1][0].startswith("'=") and rows[1][1].startswith("'@")
    xlsx_response=client.get('/api/export',headers=headers,params={'fields':'full_name,note','format':'xlsx'})
    wb=load_workbook(io.BytesIO(xlsx_response.content),data_only=False)
    assert wb.active['A2'].value.startswith("'=") and wb.active['A2'].data_type=='s'
    assert wb.active['B2'].value.startswith("'@")
    wb.close()
    json_response=client.get('/api/export',headers=headers,params={'fields':'full_name,note','format':'json'})
    assert json_response.json()[0]['full_name'].startswith('=')


def workbook_bytes(sheets):
    from openpyxl import Workbook
    wb=Workbook()
    wb.remove(wb.active)
    for title,rows in sheets:
        ws=wb.create_sheet(title)
        for row in rows:ws.append(row)
    output=io.BytesIO();wb.save(output);wb.close()
    return output.getvalue()

def test_excel_multiple_sheets_same_structure_are_merged(api):
    client,headers=api
    raw=workbook_bytes([
        ('Tổ 1',[['Họ tên','Số CCCD','Tổ'],['Nguyễn Văn An','079000000001',1]]),
        ('Tổ 2',[['Họ tên','Số CCCD','Tổ'],['Trần Thị Bình','079000000002',2]])])
    preview=client.post('/api/import/preview',headers=headers,files={'file':('nhieu-sheet.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert preview.status_code==200,preview.text
    assert preview.json()['count']==2

def test_excel_rejects_mismatched_sheets_and_formulas(api):
    client,headers=api
    mismatch=workbook_bytes([
        ('A',[['Họ tên','Tổ'],['Nguyễn A',1]]),
        ('B',[['Tên khác','Tổ'],['Nguyễn B',2]])])
    response=client.post('/api/import/preview',headers=headers,files={'file':('lech.xlsx',mismatch,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==422 and 'cùng tiêu đề' in response.json()['detail']
    formula=workbook_bytes([('A',[['Họ tên','Tổ'],['=CONCAT("Nguyễn"," A")',1]])])
    response=client.post('/api/import/preview',headers=headers,files={'file':('cong-thuc.xlsx',formula,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==422 and 'công thức' in response.json()['detail']

def test_excel_skips_hidden_sheet_and_rejects_merged_or_late_header(api):
    client,headers=api
    from openpyxl import Workbook
    wb=Workbook();visible=wb.active;visible.title='Dữ liệu'
    visible.append(['Họ tên','Tổ']);visible.append(['Nguyễn A',1])
    hidden=wb.create_sheet('Danh mục ẩn');hidden.append(['Khác','Cấu hình']);hidden.sheet_state='hidden'
    output=io.BytesIO();wb.save(output);wb.close()
    response=client.post('/api/import/preview',headers=headers,files={'file':('an.xlsx',output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==200 and response.json()['count']==1

    wb=Workbook();ws=wb.active;ws.append(['Họ tên','Tổ']);ws.append(['Nguyễn A',1]);ws.merge_cells('A1:B1')
    output=io.BytesIO();wb.save(output);wb.close()
    response=client.post('/api/import/preview',headers=headers,files={'file':('gop.xlsx',output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==422 and 'ô gộp' in response.json()['detail']

    raw=workbook_bytes([('Dữ liệu',[[None,None],['Họ tên','Tổ'],['Nguyễn A',1]])])
    response=client.post('/api/import/preview',headers=headers,files={'file':('tieu-de-muon.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==422 and 'dòng đầu tiên' in response.json()['detail']

def test_possible_duplicate_without_cccd_is_warning_not_auto_merge(api):
    client,headers=api
    payload={'full_name':'Nguyễn Thị Ánh','dob':'1990-02-03','team_id':3}
    assert client.post('/api/residents',headers=headers,json=payload).status_code==200
    check=client.post('/api/residents/duplicate-check',headers=headers,json={
        'full_name':'  NGUYEN thi anh ','dob':'1990-02-03','team_id':3})
    assert check.status_code==200 and check.json()['possible_duplicate'] is True
    commit=client.post('/api/import/commit',headers=headers,json={
        'rows':[{'Họ tên':'Nguyễn Thị Ánh','Ngày sinh':'03/02/1990','Tổ':'3'}],
        'mapping':{'Họ tên':'full_name','Ngày sinh':'dob','Tổ':'team_id'},
        'selected':['full_name','dob','team_id']})
    assert commit.status_code==200,commit.text
    assert commit.json()['inserted']==1 and commit.json()['possible_duplicates']==1
    assert len(commit.json()['warnings'])==1

def test_import_commit_rejects_duplicate_destination_mapping(api):
    client,headers=api
    response=client.post('/api/import/commit',headers=headers,json={
        'rows':[{'Tên 1':'Nguyễn A','Tên 2':'Nguyễn B','Tổ':'1'}],
        'mapping':{'Tên 1':'full_name','Tên 2':'full_name','Tổ':'team_id'},
        'selected':['full_name','team_id']})
    assert response.status_code==422 and 'cùng một trường đích' in response.json()['detail']

def test_custom_field_create_import_list_and_export(api):
    client,headers=api
    created=client.post('/api/fields',headers=headers,json={
        'name':'nhom_mau','label':'Nhóm máu','type':'dropdown','options':['A','B','AB','O']})
    assert created.status_code==200,created.text
    raw='Họ tên,Tổ,Nhóm máu\nNguyễn Văn C,4,AB\n'.encode('utf-8')
    preview=client.post('/api/import/preview',headers=headers,files={'file':('tuy-bien.csv',raw,'text/csv')})
    assert preview.status_code==200,preview.text
    data=preview.json();assert data['mapping']['Nhóm máu']=='custom:nhom_mau'
    commit=client.post('/api/import/commit',headers=headers,json={
        'rows':data['rows'],'mapping':data['mapping'],
        'selected':['full_name','team_id','custom:nhom_mau']})
    assert commit.status_code==200,commit.text
    assert commit.json()['inserted']==1
    listed=client.get('/api/residents',headers=headers).json()['items']
    assert listed[0]['custom_values']['nhom_mau']=='AB'
    exported=client.get('/api/export',headers=headers,params={
        'fields':'full_name,custom:nhom_mau','format':'json'})
    assert exported.status_code==200 and exported.json()[0]['custom:nhom_mau']=='AB'

def test_custom_field_validation_and_fail_closed_viewer_default(api):
    client,headers=api
    assert client.post('/api/fields',headers=headers,json={
        'name':'ngay_dang_ky','label':'Ngày đăng ký','type':'date'}).status_code==200
    bad=client.post('/api/residents',headers=headers,json={
        'full_name':'Nguyễn Sai Ngày','team_id':1,
        'custom_values':{'ngay_dang_ky':'2026.12.31'}})
    assert bad.status_code==422 and 'DD-MM-YYYY' in bad.json()['detail']
    matrix=client.get('/api/permissions',headers=headers).json()
    assert matrix['field_grants']['viewer']['ngay_dang_ky']=='hidden'

def test_dates_use_day_month_year_for_input_display_and_export(api):
    client,headers=api
    assert client.post('/api/fields',headers=headers,json={
        'name':'ngay_vao_o','label':'Ngày vào ở','type':'date'}).status_code==200
    created=client.post('/api/residents',headers=headers,json={
        'full_name':'Nguyễn Đúng Ngày','dob':'03-02-1990','team_id':2,
        'custom_values':{'ngay_vao_o':'27-10-2026'}})
    assert created.status_code==200,created.text
    listed=client.get('/api/residents',headers=headers).json()['items'][0]
    assert listed['dob']=='03-02-1990'
    assert listed['custom_values']['ngay_vao_o']=='27-10-2026'
    exported=client.get('/api/export',headers=headers,params={
        'fields':'full_name,dob,custom:ngay_vao_o','format':'json'})
    assert exported.status_code==200
    assert exported.json()[0]['dob']=='03-02-1990'
    assert exported.json()[0]['custom:ngay_vao_o']=='27-10-2026'

def test_legacy_iso_date_remains_readable_but_is_presented_dd_mm_yyyy(api):
    client,headers=api
    created=client.post('/api/residents',headers=headers,json={
        'full_name':'Hồ Sơ Cũ','dob':'1995-12-31','team_id':1})
    assert created.status_code==200
    assert client.get('/api/residents',headers=headers).json()['items'][0]['dob']=='31-12-1995'

def test_dropdown_requires_options_and_field_can_be_reordered_renamed_hidden(api):
    client,headers=api
    missing=client.post('/api/fields',headers=headers,json={
        'name':'phan_loai','label':'Phân loại','type':'dropdown','options':[]})
    assert missing.status_code==422
    created=client.post('/api/fields',headers=headers,json={
        'name':'phan_loai','label':'Phân loại','type':'dropdown',
        'options':['Ưu tiên','Thường','Ưu tiên']})
    assert created.status_code==200,created.text
    field_id=created.json()['id']
    updated=client.patch('/api/fields/'+field_id,headers=headers,json={
        'label':'Nhóm hồ sơ','options':['Thường','Ưu tiên'],'display_order':2,'active':False})
    assert updated.status_code==200,updated.text
    managed=client.get('/api/fields?include_inactive=true',headers=headers).json()['custom'][0]
    assert managed['label']=='Nhóm hồ sơ' and managed['display_order']==2 and managed['active']==0
    assert json.loads(managed['options_json'])==['Thường','Ưu tiên']
    assert client.get('/api/fields',headers=headers).json()['custom']==[]
    rejected=client.post('/api/residents',headers=headers,json={
        'full_name':'Nguyễn Trường Ẩn','team_id':1,'custom_values':{'phan_loai':'Thường'}})
    assert rejected.status_code==422

def test_non_dropdown_rejects_configured_options(api):
    client,headers=api
    response=client.post('/api/fields',headers=headers,json={
        'name':'ghi_chu_phu','label':'Ghi chú phụ','type':'text','options':['Sai']})
    assert response.status_code==422 and 'Chỉ trường danh sách' in response.json()['detail']

def test_dropdown_option_in_use_cannot_be_removed(api):
    client,headers=api
    created=client.post('/api/fields',headers=headers,json={
        'name':'muc_uu_tien','label':'Mức ưu tiên','type':'dropdown',
        'options':['Cao','Thường']})
    field_id=created.json()['id']
    resident=client.post('/api/residents',headers=headers,json={
        'full_name':'Nguyễn Đang Dùng','team_id':1,
        'custom_values':{'muc_uu_tien':'Cao'}})
    assert resident.status_code==200,resident.text
    blocked=client.patch('/api/fields/'+field_id,headers=headers,json={'options':['Thường']})
    assert blocked.status_code==409 and 'Cao (1 hồ sơ)' in blocked.json()['detail']
    managed=client.get('/api/fields?include_inactive=true',headers=headers).json()['custom'][0]
    assert managed['usage_count']==1 and json.loads(managed['options_json'])==['Cao','Thường']

def test_custom_fields_can_move_up_and_down_without_deleting_values(api):
    client,headers=api
    first=client.post('/api/fields',headers=headers,json={
        'name':'truong_mot','label':'Trường một','type':'text'}).json()['id']
    second=client.post('/api/fields',headers=headers,json={
        'name':'truong_hai','label':'Trường hai','type':'text'}).json()['id']
    before=client.get('/api/fields?include_inactive=true',headers=headers).json()['custom']
    assert [x['id'] for x in before]==[first,second]
    moved=client.post('/api/fields/'+second+'/move',headers=headers,json={'direction':'up'})
    assert moved.status_code==200 and moved.json()['moved'] is True
    after=client.get('/api/fields?include_inactive=true',headers=headers).json()['custom']
    assert [x['id'] for x in after]==[second,first]
    edge=client.post('/api/fields/'+second+'/move',headers=headers,json={'direction':'up'})
    assert edge.status_code==200 and edge.json()['moved'] is False

def test_dropdown_option_migration_preview_commit_and_audit(api):
    client,headers=api
    created=client.post('/api/fields',headers=headers,json={
        'name':'tinh_trang','label':'Tình trạng','type':'dropdown',
        'options':['Cũ','Mới']})
    field_id=created.json()['id']
    for index in range(2):
        assert client.post('/api/residents',headers=headers,json={
            'full_name':'Người cần chuyển '+str(index),'team_id':1,
            'custom_values':{'tinh_trang':'Cũ'}}).status_code==200
    payload={'from_value':'Cũ','to_value':'Đã xác minh'}
    preview=client.post('/api/fields/'+field_id+'/option-migration/preview',headers=headers,json=payload)
    assert preview.status_code==200 and preview.json()['affected_count']==2
    bad=client.post('/api/fields/'+field_id+'/option-migration/commit',headers=headers,json={
        **payload,'expected_count':2,'confirmation':'SAI'})
    assert bad.status_code==422
    done=client.post('/api/fields/'+field_id+'/option-migration/commit',headers=headers,json={
        **payload,'expected_count':2,'confirmation':'CHUYEN DU LIEU'})
    assert done.status_code==200 and done.json()['updated']==2
    assert done.json()['options']==['Đã xác minh','Mới']
    residents=client.get('/api/residents',headers=headers).json()['items']
    assert all(x['custom_values']['tinh_trang']=='Đã xác minh' for x in residents)
    audit=client.get('/api/audit',headers=headers).json()['items']
    assert any(x['action']=='MIGRATE_FIELD_OPTION' for x in audit)

def test_dropdown_migration_stops_when_count_changed_after_preview(api):
    client,headers=api
    field_id=client.post('/api/fields',headers=headers,json={
        'name':'nhom_xu_ly','label':'Nhóm xử lý','type':'dropdown',
        'options':['A','B']}).json()['id']
    payload={'from_value':'A','to_value':'B'}
    assert client.post('/api/residents',headers=headers,json={
        'full_name':'Người thứ nhất','team_id':1,'custom_values':{'nhom_xu_ly':'A'}}).status_code==200
    preview=client.post('/api/fields/'+field_id+'/option-migration/preview',headers=headers,json=payload).json()
    assert client.post('/api/residents',headers=headers,json={
        'full_name':'Người thứ hai','team_id':1,'custom_values':{'nhom_xu_ly':'A'}}).status_code==200
    stopped=client.post('/api/fields/'+field_id+'/option-migration/commit',headers=headers,json={
        **payload,'expected_count':preview['affected_count'],'confirmation':'CHUYEN DU LIEU'})
    assert stopped.status_code==409 and 'xem trước lại' in stopped.json()['detail']

def test_dropdown_migration_preview_details_history_and_safe_undo(api):
    client,headers=api
    field_id=client.post('/api/fields',headers=headers,json={
        'name':'ket_qua','label':'Kết quả','type':'dropdown','options':['Chờ','Xong']}).json()['id']
    resident=client.post('/api/residents',headers=headers,json={
        'full_name':'Nguyễn Cần Xem','dob':'28-10-1990','team_id':3,
        'custom_values':{'ket_qua':'Chờ'}})
    assert resident.status_code==200
    payload={'from_value':'Chờ','to_value':'Xong'}
    preview=client.post('/api/fields/'+field_id+'/option-migration/preview',headers=headers,json=payload)
    assert preview.status_code==200
    assert preview.json()['sample'][0]['full_name']=='Nguyễn Cần Xem'
    assert preview.json()['sample'][0]['dob']=='28-10-1990'
    migrated=client.post('/api/fields/'+field_id+'/option-migration/commit',headers=headers,json={
        **payload,'expected_count':1,'confirmation':'CHUYEN DU LIEU'})
    migration_id=migrated.json()['migration_id']
    history=client.get('/api/fields/'+field_id+'/option-migrations',headers=headers)
    assert history.status_code==200 and history.json()['items'][0]['id']==migration_id
    undone=client.post('/api/fields/'+field_id+'/option-migrations/'+migration_id+'/undo',headers=headers,json={'confirmation':'HOAN TAC'})
    assert undone.status_code==200 and undone.json()['restored']==1
    listed=client.get('/api/residents',headers=headers).json()['items'][0]
    assert listed['custom_values']['ket_qua']=='Chờ'
    second=client.post('/api/fields/'+field_id+'/option-migrations/'+migration_id+'/undo',headers=headers,json={'confirmation':'HOAN TAC'})
    assert second.status_code==409 and 'đã được hoàn tác' in second.json()['detail']

def test_dropdown_undo_stops_if_migrated_resident_was_edited(api):
    client,headers=api
    field_id=client.post('/api/fields',headers=headers,json={
        'name':'trang_thai_phu','label':'Trạng thái phụ','type':'dropdown',
        'options':['Một','Hai','Ba']}).json()['id']
    created=client.post('/api/residents',headers=headers,json={
        'full_name':'Người Đã Sửa','team_id':1,'custom_values':{'trang_thai_phu':'Một'}})
    resident_id=created.json()['id'];payload={'from_value':'Một','to_value':'Hai'}
    migrated=client.post('/api/fields/'+field_id+'/option-migration/commit',headers=headers,json={
        **payload,'expected_count':1,'confirmation':'CHUYEN DU LIEU'})
    migration_id=migrated.json()['migration_id']
    edited=client.put('/api/residents/'+resident_id,headers=headers,json={
        'full_name':'Người Đã Sửa','team_id':1,'custom_values':{'trang_thai_phu':'Ba'}})
    assert edited.status_code==200
    stopped=client.post('/api/fields/'+field_id+'/option-migrations/'+migration_id+'/undo',headers=headers,json={'confirmation':'HOAN TAC'})
    assert stopped.status_code==409 and 'hồ sơ đã thay đổi' in stopped.json()['detail']

def test_duplicate_review_center_lists_comparison_and_records_not_duplicate(api):
    client,headers=api
    for phone in ('0900000001','0900000002'):
        created=client.post('/api/residents',headers=headers,json={
            'full_name':'Nguyễn Văn Trùng','dob':'01-02-1990','team_id':2,'phone':phone})
        assert created.status_code==200,created.text
    candidates=client.get('/api/residents/duplicate-candidates',headers=headers)
    assert candidates.status_code==200 and candidates.json()['count']==1
    item=candidates.json()['items'][0]
    assert item['left']['dob']=='01-02-1990' and item['right']['dob']=='01-02-1990'
    decided=client.post('/api/residents/duplicate-candidates/'+item['candidate_id']+'/decision',
                        headers=headers,json={'decision':'not_duplicate'})
    assert decided.status_code==200
    assert client.get('/api/residents/duplicate-candidates',headers=headers).json()['count']==0
    audit=client.get('/api/audit',headers=headers).json()['items']
    assert any(x['action']=='DUPLICATE_NOT_MATCH' for x in audit)

def test_duplicate_review_requires_missing_cccd_and_changes_create_new_candidate(api):
    client,headers=api
    base={'full_name':'Trần Thị Giống','dob':'05-06-1985','team_id':1}
    first=client.post('/api/residents',headers=headers,json={**base,'cccd':'079000000111'}).json()['id']
    second=client.post('/api/residents',headers=headers,json={**base,'cccd':'079000000222'}).json()['id']
    assert client.get('/api/residents/duplicate-candidates',headers=headers).json()['count']==0
    # Removing one CCCD makes the pair eligible for manual review.
    updated=client.put('/api/residents/'+second,headers=headers,json=base)
    assert updated.status_code==200
    item=client.get('/api/residents/duplicate-candidates',headers=headers).json()['items'][0]
    assert {item['left']['id'],item['right']['id']}=={first,second}
