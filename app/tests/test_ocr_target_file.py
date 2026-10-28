"""Selected OCR columns must append to exactly one matching XLSX sheet."""
import io
from openpyxl import Workbook,load_workbook
from test_foundation_accounts import api  # noqa: F401

FIELDS=[{'label':'Họ và tên','column':'Họ và tên','aliases':[],'type':'text'},
        {'label':'Số hiệu bằng','column':'Số hiệu bằng','aliases':[],'type':'identifier'}]

def workbook_bytes(headers=('Họ và tên','Số hiệu bằng'),duplicate=False):
    wb=Workbook();ws=wb.active;ws.title='Danh sách';ws.append(list(headers));ws.append(['Nguyễn A','12345/PT'])
    other=wb.create_sheet('Ghi chú');other.append(['Giữ nguyên']);other.append(['Nội dung cũ'])
    if duplicate:
        dup=wb.create_sheet('Danh sách 2');dup.append(list(headers))
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()

def test_append_reviewed_ocr_to_selected_excel_and_keep_other_sheets(api):
    client,auth,_=api
    saved=client.post('/api/ocr/templates',headers=auth,json={'name':'Bằng cấp','fields':FIELDS})
    assert saved.status_code==200,saved.text
    tid=saved.json()['id'];source=workbook_bytes()
    fields={'template_id':tid}
    check=client.post('/api/ocr/target/validate',headers=auth,data=fields,files={'file':('nguon.xlsx',source)})
    assert check.status_code==200 and check.json()=={'sheet':'Danh sách','existing_rows':1,'columns':['Họ và tên','Số hiệu bằng']}
    import json
    payload={'template_id':tid,'filename':'ket_qua','rows_json':json.dumps([{'Họ và tên':'Trần B','Số hiệu bằng':'=1+1'}])}
    output=client.post('/api/ocr/export-into-file',headers=auth,data=payload,files={'file':('nguon.xlsx',source)})
    assert output.status_code==200,output.text
    assert output.content!=source
    wb=load_workbook(io.BytesIO(output.content));ws=wb['Danh sách']
    assert [c.value for c in ws[1]]==['Họ và tên','Số hiệu bằng']
    assert [c.value for c in ws[2]]==['Nguyễn A','12345/PT']
    assert [c.value for c in ws[3]]==['Trần B',"'=1+1"]
    assert wb['Ghi chú']['A2'].value=='Nội dung cũ'
    wb.close()

def test_wrong_header_order_and_ambiguous_sheet_are_rejected(api):
    client,auth,_=api
    tid=client.post('/api/ocr/templates',headers=auth,json={'name':'Bằng cấp','fields':FIELDS}).json()['id']
    for source in (workbook_bytes(('Số hiệu bằng','Họ và tên')),workbook_bytes(duplicate=True)):
        check=client.post('/api/ocr/target/validate',headers=auth,data={'template_id':tid},files={'file':('sai.xlsx',source)})
        assert check.status_code==422 and 'đúng một sheet' in check.json()['detail']
        export=client.post('/api/ocr/export-into-file',headers=auth,data={'template_id':tid,'filename':'ket_qua','rows_json':'[{"Họ và tên":"Trần B","Số hiệu bằng":"123"}]'},files={'file':('sai.xlsx',source)})
        assert export.status_code==422

def test_export_rejects_unconfigured_column(api):
    client,auth,_=api
    tid=client.post('/api/ocr/templates',headers=auth,json={'name':'Bằng cấp','fields':FIELDS}).json()['id']
    result=client.post('/api/ocr/export-into-file',headers=auth,data={'template_id':tid,'filename':'ket_qua','rows_json':'[{"Họ và tên":"Trần B","Nghề nghiệp":"Kế toán"}]'},files={'file':('nguon.xlsx',workbook_bytes())})
    assert result.status_code==422 and 'không khớp' in result.json()['detail']
