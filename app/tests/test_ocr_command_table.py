"""A scanned header row defines the output; selected fields alone receive OCR/QR."""
import io
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from test_foundation_accounts import api  # noqa: F401

HEADERS=['Mã hồ sơ','Họ và tên','Năm sinh','Địa chỉ','Nghề nghiệp','Tổ','Số điện thoại','Ghi chú','Ngày tạo','Nguồn']
SELECTED={'Họ và tên','Năm sinh','Địa chỉ','Nghề nghiệp'}

def target_file():
    wb=Workbook();ws=wb.active;ws.title='Danh sách';ws.append(HEADERS)
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

def fields():
    return [{'label':name,'column':name,'aliases':[],'type':'year' if name=='Năm sinh' else 'text',
             'extract':name in SELECTED} for name in HEADERS]

def test_scan_headers_build_command_and_export_only_requested_values(api):
    client,headers,_=api
    source=target_file()
    scanned=client.post('/api/ocr/columns/scan',headers=headers,files={'file':('cot.xlsx',source)})
    assert scanned.status_code==200,scanned.text
    assert scanned.json()['sheets']==[{'name':'Danh sách','columns':HEADERS}]
    created=client.post('/api/ocr/templates',headers=headers,json={'name':'Bảng lệnh 10 cột','fields':fields()})
    assert created.status_code==200,created.text
    tid=created.json()['id']
    text='Họ và tên: Nguyễn Văn A\nNăm sinh: 1988\nĐịa chỉ: Long An\nNghề nghiệp: Kế toán\nSố điện thoại: 0900000000\nGhi chú: không nhập'
    extracted=client.post('/api/ocr/extract',headers=headers,json={'template_id':tid,'text':text})
    assert extracted.status_code==200,extracted.text
    row=extracted.json()['values']
    assert list(row)==HEADERS
    assert {key:row[key] for key in SELECTED}=={'Họ và tên':'Nguyễn Văn A','Năm sinh':'1988','Địa chỉ':'Long An','Nghề nghiệp':'Kế toán'}
    assert all(row[key]=='' for key in HEADERS if key not in SELECTED)
    export=client.post('/api/ocr/export',headers=headers,json={'template_id':tid,'filename':'ket_qua','rows':[row]})
    assert export.status_code==200,export.text
    wb=load_workbook(io.BytesIO(export.content));ws=wb.active
    assert [cell.value for cell in ws[1]]==HEADERS
    assert [cell.value or '' for cell in ws[2]]==[row[h] for h in HEADERS]
    wb.close()
    bad=client.post('/api/ocr/export',headers=headers,json={'template_id':tid,'filename':'ket_qua','rows':[{**row,'Số điện thoại':'0900000000'}]})
    assert bad.status_code==422
    with patch('server.app.decode_qr_file',return_value=[{'page':1,'text':'Họ và tên: Nguyễn Văn B\nSố điện thoại: 0900000000'}]):
        qr=client.post('/api/qr/decode?template_id='+tid,headers=headers,files={'file':('scan.png',b'content','image/png')})
    assert qr.status_code==200,qr.text
    assert qr.json()['values']['Họ và tên']=='Nguyễn Văn B'
    assert qr.json()['values']['Số điện thoại']==''
    with patch('server.app.parse_file',return_value={'headers':[],'rows':[],'mapping':{},'count':0}) as parse:
        preview=client.post('/api/import/preview?ocr_template_id='+tid,headers=headers,files={'file':('scan.png',b'content','image/png')})
    assert preview.status_code==200,preview.text
    assert [f['column'] for f in parse.call_args.kwargs['requested_fields']]==[h for h in HEADERS if h in SELECTED]

def test_command_requires_at_least_one_selected_field(api):
    client,headers,_=api
    inactive=[{**field,'extract':False} for field in fields()]
    response=client.post('/api/ocr/templates',headers=headers,json={'name':'Bảng trống','fields':inactive})
    assert response.status_code==422
